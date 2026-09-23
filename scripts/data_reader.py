# -*- coding: utf-8 -*-
"""
Đọc file Excel → list dict.
⚡ KHI ĐỔI CẤU TRÚC EXCEL: chỉ cần sửa hàm read_excel() bên dưới.
Format chuẩn trả về: [{"stt","hsk","topic","subject","vi","zh","pinyin"}, ...]
"""
import openpyxl
import os
import sys


def clean(s):
    if s is None:
        return ""
    return (str(s).replace('\n', ' ').replace('\r', ' ')
            .replace('\t', ' ').replace('\\', '\\\\'))


def read_excel(excel_file, sheet_index=0):
    print(f"\n📖 Đang đọc file: {excel_file}")
    if not os.path.exists(excel_file):
        print(f"❌ Không tìm thấy file {excel_file}")
        sys.exit(1)

    wb = openpyxl.load_workbook(excel_file, data_only=True)
    ws = wb.worksheets[sheet_index]
    print(f"📊 Sheet: {ws.title} - {ws.max_row} dòng")

    # ↓↓↓ CẤU HÌNH CỘT Ở ĐÂY (đổi khi Excel đổi cấu trúc) ↓↓↓
    COL_STT = 0
    COL_HSK = 1
    COL_TOPIC = 2
    COL_SUBJECT = 3
    COL_VI = 4
    COL_ZH = 5
    COL_PINYIN = 6
    DATA_START = 2
    # ↑↑↑ HẾT PHẦN CẦN SỬA ↑↑↑

    data = []
    for row in ws.iter_rows(min_row=DATA_START, values_only=True):
        if not row or len(row) <= max(COL_VI, COL_ZH):
            continue

        stt = row[COL_STT] if COL_STT < len(row) and row[COL_STT] is not None else ""
        hsk = clean(row[COL_HSK]) if COL_HSK < len(row) else ""
        topic = clean(row[COL_TOPIC]) if COL_TOPIC < len(row) else ""
        subject = clean(row[COL_SUBJECT]) if COL_SUBJECT < len(row) else ""
        vi = clean(row[COL_VI]) if COL_VI < len(row) else ""
        zh = clean(row[COL_ZH]) if COL_ZH < len(row) else ""
        pinyin = clean(row[COL_PINYIN]) if COL_PINYIN < len(row) else ""

        if not vi and not zh:
            continue

        data.append({
            "stt": str(stt), "hsk": hsk, "topic": topic, "subject": subject,
            "vi": vi, "zh": zh, "pinyin": pinyin
        })

    print(f"✅ Đã đọc {len(data)} câu")
    return data
