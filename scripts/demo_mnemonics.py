# -*- coding: utf-8 -*-
r"""
Demo sinh mẹo nhớ cho vocab HSK bằng OpenRouter + Qwen.
Chạy trên GitHub Actions.

Input:
    - Environment: OPENROUTER_API_KEY (từ GitHub Secrets)
    - File: data/tu_vung_hsk.xlsx

Output:
    - data/mnemonics_demo.json
"""
import os
import sys
import json
import time
from openai import OpenAI
import openpyxl


# ═══════════════════════════════════════════════════════════════════
#  CONFIG
# ═══════════════════════════════════════════════════════════════════
API_KEY = os.getenv("OPENROUTER_API_KEY")

if not API_KEY:
    print("❌ Chưa set OPENROUTER_API_KEY")
    sys.exit(1)

client = OpenAI(
    api_key=API_KEY,
    base_url="https://openrouter.ai/api/v1"
)

EXCEL_FILE = "data/tu_vung_hsk.xlsx"
OUTPUT_FILE = "data/mnemonics_demo.json"

# ⭐ Model Qwen miễn phí
MODEL_ID = os.getenv("QWEN_MODEL", "qwen/qwen3.5-flash:free")

# ⭐ Số từ demo (có thể truyền qua env)
DEMO_LIMIT = int(os.getenv("DEMO_LIMIT", "20"))

# ⭐ Delay giữa requests (giây)
DELAY_BETWEEN = float(os.getenv("DELAY_BETWEEN", "2.0"))


# ═══════════════════════════════════════════════════════════════════
#  PROMPT
# ═══════════════════════════════════════════════════════════════════
def build_prompt(zh, pinyin, vi, hsk):
    return f"""Bạn là giáo viên tiếng Trung cho người Việt Nam.

Nhiệm vụ: Sinh mẹo nhớ cho từ: **{zh}** (pinyin: {pinyin}) - nghĩa: {vi} - HSK{hsk}

Yêu cầu output có 4-5 dòng, mỗi dòng bắt đầu bằng emoji:

💡 Chiết tự: [Phân tích bộ thủ, nếu từ ghép phân tích từng chữ]
📌 Âm thanh: [Liên tưởng âm Hán Việt với từ tiếng Việt quen thuộc]
🎬 Câu chuyện: [1-2 câu liên kết các bộ thủ/âm thanh thành câu chuyện dễ nhớ]
📎 Ví dụ: [câu ví dụ tiếng Trung ngắn + nghĩa Việt]
🔗 Liên quan: [2-3 từ cùng chữ/bộ/âm]

Quy tắc:
- Ngắn gọn, dễ đọc trên mobile
- KHÔNG dùng ký tự đặc biệt ngoài emoji
- Câu chuyện phải logic, dễ hình dung
- Nếu từ 1 chữ: tập trung chiết tự
- Nếu từ 2+ chữ: phân tích nghĩa từng chữ rồi ghép
- Nếu có âm gần giống tiếng Việt: BẮT BUỘC dùng để liên tưởng

Ví dụ cho 菜 (cài - món ăn):
💡 Chiết tự: 艹 (cỏ) + 采 (hái) → Hái cỏ về nấu = MÓN ĂN
📌 Âm thanh: "cài" ≈ "cải" (rau cải) - cùng họ rau củ
🎬 Câu chuyện: Đi hái (采) rau cỏ (艹) → chế biến thành món ăn (菜)
📎 Ví dụ: 这道菜很好吃 (Món này rất ngon)
🔗 Liên quan: 采 (cǎi - hái), 彩 (cǎi - màu sắc)

Bây giờ sinh mẹo cho: {zh}
Output:"""


# ═══════════════════════════════════════════════════════════════════
#  GỌI API
# ═══════════════════════════════════════════════════════════════════
def generate_mnemonic(zh, pinyin, vi, hsk):
    prompt = build_prompt(zh, pinyin, vi, hsk)
    
    for attempt in range(3):
        try:
            response = client.chat.completions.create(
                model=MODEL_ID,
                messages=[{"role": "user", "content": prompt}],
                temperature=0.7,
                max_tokens=500,
                extra_headers={
                    "HTTP-Referer": "https://github.com",
                    "X-Title": "Vocab Mnemonic Generator"
                }
            )
            content = response.choices[0].message.content.strip()
            if content:
                return content
        except Exception as e:
            err = str(e)
            if "429" in err or "rate" in err.lower():
                wait = 15 * (attempt + 1)
                print(f"      ⏳ Rate limit, chờ {wait}s...")
                time.sleep(wait)
            else:
                print(f"      ⚠️  Lỗi (lần {attempt+1}/3): {err[:150]}")
                time.sleep(3)
    return ""


# ═══════════════════════════════════════════════════════════════════
#  LOAD EXCEL
# ═══════════════════════════════════════════════════════════════════
def load_sample_words():
    if not os.path.exists(EXCEL_FILE):
        print(f"❌ Không tìm thấy {EXCEL_FILE}")
        return []
    
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True, read_only=True)
    samples = []
    
    for sheet_name in wb.sheetnames:
        if len(samples) >= DEMO_LIMIT:
            break
        if not sheet_name.strip().upper().startswith("HSK"):
            continue
        
        ws = wb[sheet_name]
        is_hsk79 = "7-9" in sheet_name
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(samples) >= DEMO_LIMIT:
                break
            if not row or len(row) < 3:
                continue
            
            stt = str(row[0] or "").strip()
            zh = str(row[1] or "").strip()
            pinyin = str(row[2] or "").strip()
            
            vi_idx = 4 if is_hsk79 else 5
            vi = str(row[vi_idx] or "").strip() if len(row) > vi_idx else ""
            
            if not zh or not vi:
                continue
            if len(zh) > 3:  # Chỉ lấy từ ngắn
                continue
            
            samples.append({
                "stt": stt,
                "zh": zh,
                "pinyin": pinyin,
                "vi": vi,
                "hsk": sheet_name.strip(),
                "key": f"{sheet_name.strip()}|{stt}|{zh}"
            })
    
    wb.close()
    return samples[:DEMO_LIMIT]


# ═══════════════════════════════════════════════════════════════════
#  MAIN
# ═══════════════════════════════════════════════════════════════════
def main():
    print("=" * 62)
    print("🎬 DEMO SINH MẸO NHỚ — OPENROUTER + QWEN")
    print(f"   Model: {MODEL_ID}")
    print(f"   Số từ: {DEMO_LIMIT}")
    print(f"   Delay: {DELAY_BETWEEN}s")
    print("=" * 62)
    
    samples = load_sample_words()
    print(f"\n📚 Đã lấy {len(samples)} từ mẫu\n")
    
    if not samples:
        print("❌ Không có từ nào để test")
        sys.exit(1)
    
    results = {}
    success = 0
    failed = 0
    
    for i, w in enumerate(samples, 1):
        preview = w['vi'][:40] + ('...' if len(w['vi']) > 40 else '')
        print(f"[{i}/{len(samples)}] {w['zh']} ({w['pinyin']}) - {preview}")
        
        mnemonic = generate_mnemonic(w["zh"], w["pinyin"], w["vi"], w["hsk"])
        
        if mnemonic:
            results[w["key"]] = mnemonic
            success += 1
            print(f"   ✅ OK")
        else:
            failed += 1
            print(f"   ❌ FAIL")
        
        # Lưu sau mỗi từ
        os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
        with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
            json.dump(results, f, ensure_ascii=False, indent=2)
        
        time.sleep(DELAY_BETWEEN)
    
    print("\n" + "=" * 62)
    print(f"✅ HOÀN TẤT: {success}/{len(samples)} từ")
    print(f"📁 File: {OUTPUT_FILE}")
    print("=" * 62)
    
    # In mẫu
    print("\n📝 3 MẪU ĐẦU TIÊN:\n")
    for i, (key, mnemonic) in enumerate(list(results.items())[:3], 1):
        zh = key.split('|')[-1]
        print(f"━━━ Từ {i}: {zh} ━━━")
        print(mnemonic)
        print()


if __name__ == "__main__":
    main()
