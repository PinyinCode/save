# -*- coding: utf-8 -*-
"""
fix.py — Auto-scan data/ và thêm MỌI file Excel thành tab riêng.
+ TỰ ĐỘNG thêm tab TỪ VỰNG PREMIUM (👑) từ data/tu_vung_hsk.xlsx

ĐẶC ĐIỂM:
  - Đọc HẾT mọi file .xlsx/.xls/.csv trong data/ (TRỪ input.xlsx)
  - Tên tab = TÊN FILE (normalize NFC — hiển thị đúng dấu tiếng Việt)
  - Label tab = CHỈ tên file, KHÔNG thêm số câu
  - Logic đọc Excel GIỐNG data_reader.py
  - TIER LOCK + ONBOARDING giống tab tổng hợp
  - Clone nút để XÓA event listener cũ → không còn popup
  - CHỈ 1 TAB ACTIVE tại một thời điểm
  - KHÔNG inject vào DATASET_REGISTRY
  - KHÔNG can thiệp convert.py, ui_template.py, config.json

  ⭐ TỪ VỰNG PREMIUM (mới):
  - Đọc file data/tu_vung_hsk.xlsx (11 cột A-K)
  - Tự sinh mẹo nhớ + bộ thủ từ vocab_data/
  - CHỈ mở cho Admin + Premium (isPermanent)
  - Modal upgrade khi không có quyền

Cách chạy:
    python scripts/convert.py    # Tạo index.html gốc
    python fix.py                # Patch index.html
"""
import json
import os
import re
import sys
import glob
import unicodedata

import openpyxl


# =================================================================
#  CONFIG
# =================================================================
INDEX_HTML = "index.html"
CONFIG_JSON = "config.json"
DATA_DIR = "data"

# File cần BỎ QUA khi quét data/ (vì đã là tab "Tổng hợp" trong convert.py)
SKIP_FILES = {"input.xlsx", "input.xls", "input.csv"}

# ⭐ File TỪ VỰNG PREMIUM (đặt cùng cấp input.xlsx)
VOCAB_FILE = os.path.join(DATA_DIR, "tu_vung_hsk.xlsx")
VOCAB_ID = "tu-vung"
VOCAB_LABEL = "Từ vựng HSK"

if not os.path.isfile(CONFIG_JSON) and os.path.isfile(os.path.join("..", CONFIG_JSON)):
    os.chdir("..")
    print("[fix.py] Phat hien chay tu scripts/ -> chuyen ve root")

TAB_ICONS = [
    "fa-comments", "fa-file-alt", "fa-book", "fa-graduation-cap",
    "fa-star", "fa-fire", "fa-bolt", "fa-rocket",
]
TAB_COLORS = [
    "#0891b2", "#dc2626", "#059669", "#d97706",
    "#7c3aed", "#db2777", "#0284c7", "#65a30d",
]


# ═══════════════════════════════════════════════════════════════════
#  ⭐ IMPORT MODULE TỪ VỰNG PREMIUM (có try/except để không crash)
# ═══════════════════════════════════════════════════════════════════
try:
    from vocab_premium import (
        read_vocab_excel,
        build_vocab_css,
        build_vocab_tab_html,
        build_vocab_modal_html,
        build_vocab_js_override,
    )
    HAS_VOCAB_MODULE = True
    print("[fix.py] ✅ Đã load module vocab_premium")
except ImportError as e:
    HAS_VOCAB_MODULE = False
    print("[fix.py] ⚠️  Không load được vocab_premium: " + str(e))
    print("[fix.py] → Tab Từ vựng Premium sẽ KHÔNG được thêm")


# =================================================================
#  CHUYEN SO A RAP -> SO HAN
# =================================================================
try:
    import cn2an
    HAS_CN2AN = True
except ImportError:
    HAS_CN2AN = False
    print("[fix.py] Khong co cn2an - dung bo dich so du phong")

_CN_DIGITS = ['零', '一', '二', '三', '四', '五', '六', '七', '八', '九']
_CN_UNITS = ['', '十', '百', '千']


def _num_to_chinese_basic(num):
    if num == 0:
        return '零'
    if num < 0:
        return '负' + _num_to_chinese_basic(-num)
    result = ''
    unit_idx = 0
    n = num
    while n > 0:
        digit = n % 10
        if digit != 0:
            if not (unit_idx == 1 and digit == 1 and n < 20 and result == ''):
                result = _CN_DIGITS[digit] + _CN_UNITS[unit_idx] + result
            else:
                result = _CN_UNITS[unit_idx] + result
        else:
            if result and not result.startswith('零'):
                result = '零' + result
        n //= 10
        unit_idx += 1
    if result.startswith('一十'):
        result = result[1:]
    return result


def _num_to_chinese(num):
    if HAS_CN2AN:
        try:
            return cn2an.an2cn(num)
        except Exception:
            pass
    return _num_to_chinese_basic(num)


def _digits_to_chinese(digits_str):
    return ''.join([_CN_DIGITS[int(d)] for d in digits_str])


def convert_arabic_to_chinese(text):
    if not text or not isinstance(text, str):
        return text

    def replace_percent(match):
        num_str = match.group(1)
        if '.' in num_str:
            parts = num_str.split('.')
            int_part = _num_to_chinese(int(parts[0]))
            dec_part = _digits_to_chinese(parts[1])
            return '百分之' + int_part + '点' + dec_part
        return '百分之' + _num_to_chinese(int(num_str))

    text = re.sub(r'(\d+(?:\.\d+)?)%', replace_percent, text)

    def replace_decimal(match):
        num_str = match.group(0)
        parts = num_str.split('.')
        int_part = _num_to_chinese(int(parts[0]))
        dec_part = _digits_to_chinese(parts[1])
        return int_part + '点' + dec_part

    text = re.sub(
        r'(?<![A-Za-z\-\.])\d+\.\d+(?![A-Za-z])',
        replace_decimal, text
    )

    def replace_int(match):
        return _num_to_chinese(int(match.group(0)))

    text = re.sub(
        r'(?<![A-Za-z\-\.])\d+(?![A-Za-z]|\.\d)',
        replace_int, text
    )

    return text


def _clean(s):
    if s is None:
        return ""
    return (str(s).replace('\n', ' ').replace('\r', ' ')
            .replace('\t', ' ').replace('\\', '\\\\'))


# =================================================================
#  HELPERS
# =================================================================
def _slugify(filename):
    base = filename.rsplit(".", 1)[0]
    base = unicodedata.normalize("NFD", base)
    base = "".join(c for c in base if unicodedata.category(c) != "Mn")
    base = base.replace("đ", "d").replace("Đ", "D")
    return re.sub(r"[^a-zA-Z0-9]+", "-", base).strip("-").lower() or "dataset"


def _display_name(filename):
    name = filename.rsplit(".", 1)[0].replace("_", " ").strip()
    name = unicodedata.normalize("NFC", name)
    if name.islower() or name.isupper():
        name = name.title()
    return name


def _escape_json_for_script(obj):
    s = json.dumps(obj, ensure_ascii=False, separators=(",", ":"))
    return s.replace("</", "<\\/")


def _js_str(s):
    if s is None:
        return ""
    return (str(s)
            .replace("\\", "\\\\")
            .replace('"', '\\"')
            .replace("'", "\\'")
            .replace("\n", "\\n")
            .replace("\r", "\\r")
            .replace("</", "<\\/"))


def _find_file_safe(filepath):
    if os.path.isfile(filepath):
        return filepath
    dirname = os.path.dirname(filepath) or "."
    basename = os.path.basename(filepath)
    if not os.path.isdir(dirname):
        return None
    variants = set()
    variants.add(basename)
    variants.add(unicodedata.normalize("NFC", basename))
    variants.add(unicodedata.normalize("NFD", basename))
    try:
        for fname in os.listdir(dirname):
            fname_nfc = unicodedata.normalize("NFC", fname)
            fname_nfd = unicodedata.normalize("NFD", fname)
            for variant in variants:
                if (fname == variant
                        or fname_nfc == unicodedata.normalize("NFC", variant)
                        or fname_nfd == unicodedata.normalize("NFD", variant)):
                    return os.path.join(dirname, fname)
    except Exception:
        pass
    return None


def _read_excel_rows(filepath):
    real_path = _find_file_safe(filepath)
    if not real_path:
        print("      [X] Khong tim thay file: " + os.path.basename(filepath))
        return []

    try:
        wb = openpyxl.load_workbook(real_path, data_only=True)
    except Exception as e:
        print("      [X] Loi load: " + type(e).__name__ + ": " + str(e))
        return []

    try:
        ws = wb.worksheets[0]
    except Exception as e:
        print("      [X] Loi sheet: " + str(e))
        return []

    print("      Sheet: " + ws.title
          + " - " + str(ws.max_row) + " dong, "
          + str(ws.max_column) + " cot")

    COL_STT = 0
    COL_HSK = 1
    COL_TOPIC = 2
    COL_SUBJECT = 3
    COL_VI = 4
    COL_ZH = 5
    COL_PINYIN = 6
    DATA_START = 2

    rows = []
    converted_count = 0
    skipped_empty = 0

    for row in ws.iter_rows(min_row=DATA_START, values_only=True):
        if not row or len(row) <= max(COL_VI, COL_ZH):
            skipped_empty += 1
            continue

        stt_val = row[COL_STT] if COL_STT < len(row) and row[COL_STT] is not None else ""
        hsk = _clean(row[COL_HSK]) if COL_HSK < len(row) else ""
        topic = _clean(row[COL_TOPIC]) if COL_TOPIC < len(row) else ""
        subject = _clean(row[COL_SUBJECT]) if COL_SUBJECT < len(row) else ""
        vi = _clean(row[COL_VI]) if COL_VI < len(row) else ""
        zh = _clean(row[COL_ZH]) if COL_ZH < len(row) else ""
        pinyin = _clean(row[COL_PINYIN]) if COL_PINYIN < len(row) else ""

        if not vi and not zh:
            skipped_empty += 1
            continue

        zh_original = zh
        zh = convert_arabic_to_chinese(zh)
        if zh != zh_original:
            converted_count += 1

        rows.append({
            "stt": str(stt_val),
            "hsk": hsk,
            "topic": topic,
            "subject": subject,
            "vi": vi,
            "zh": zh,
            "pinyin": pinyin,
        })

    print("      [OK] " + str(len(rows)) + " cau (bo qua "
          + str(skipped_empty) + " dong rong)")
    if converted_count > 0:
        print("      [OK] Chuyen so A Rap -> Han: " + str(converted_count) + " cau")

    return rows


# =================================================================
#  SCAN data/ — BO QUA input.xlsx VÀ tu_vung_hsk.xlsx
# =================================================================
def scan_data_dir():
    if not os.path.isdir(DATA_DIR):
        print("[fix.py] Khong thay thu muc '" + DATA_DIR + "/' - bo qua.")
        return []

    files = []
    for ext in ("*.xlsx", "*.xls", "*.csv"):
        files.extend(glob.glob(os.path.join(DATA_DIR, ext)))
    files = sorted(set(files))

    if not files:
        print("[fix.py] Khong co file Excel trong '" + DATA_DIR + "/'")
        return []

    print("")
    print("[fix.py] Quet '" + DATA_DIR + "/' - " + str(len(files)) + " file")

    datasets = []
    used_ids = set()

    # ⭐ Tên file từ vựng để bỏ qua khi quét tab thường
    vocab_basename = os.path.basename(VOCAB_FILE).lower()

    for filepath in files:
        fname = os.path.basename(filepath)

        if fname.startswith("~$"):
            print("   [skip] " + fname + " - file tam")
            continue

        if fname.lower() in SKIP_FILES:
            print("   [skip] " + fname + " - da la tab Tong hop")
            continue

        # ⭐ BỎ QUA file từ vựng (sẽ tạo tab Premium riêng)
        if fname.lower() == vocab_basename:
            print("   [skip] " + fname + " - se tao tab Tu vung Premium rieng")
            continue

        print("   [file] " + fname)
        rows = _read_excel_rows(filepath)
        if not rows:
            print("   [!] " + fname + " - rong hoac loi, bo qua")
            continue

        base_id = _slugify(fname)
        dataset_id = base_id
        counter = 2
        while dataset_id in used_ids:
            dataset_id = base_id + "-" + str(counter)
            counter += 1
        used_ids.add(dataset_id)

        idx = len(datasets)
        icon = TAB_ICONS[idx % len(TAB_ICONS)]
        color = TAB_COLORS[idx % len(TAB_COLORS)]

        display = _display_name(fname)
        datasets.append({
            "id": dataset_id,
            "name": display,
            "icon": icon,
            "color": color,
            "data": rows,
            "count": len(rows),
            "source": fname,
            "type": "main",
            "group": "fixpy",
        })
        print("   [OK] " + fname + " -> tab '" + display
              + "' (" + str(len(rows)) + " cau)")

    return datasets


# =================================================================
#  BUILD JS OVERRIDE
# =================================================================
def build_js_override(ids_js, datasets_json):
    """Tra ve chuoi JS override - co onboarding + tier lock."""
    L = []
    add = L.append

    add("")
    add("<script>")
    add("/* ================================================================")
    add("   FIX.PY OVERRIDE - Bind tab moi + TIER LOCK + ONBOARDING")
    add("   ================================================================")
    add("*/")
    add("(function() {")
    add("    'use strict';")
    add("    var NEW_IDS = " + ids_js + ";")
    add("")

    # FIXPY_DATASETS
    add("    window.FIXPY_DATASETS = " + datasets_json + ";")
    add("")

    # Patch __switchRawData
    add("    function patchSwitchRawData() {")
    add("        if (window.__fixPySwitchPatched) return;")
    add("        var origSwitch = window.__switchRawData;")
    add("        if (typeof origSwitch !== 'function') {")
    add("            setTimeout(patchSwitchRawData, 100);")
    add("            return;")
    add("        }")
    add("        window.__switchRawData = function(datasetId) {")
    add("            if (window.FIXPY_DATASETS && window.FIXPY_DATASETS[datasetId]) {")
    add("                RAW_DATA = window.FIXPY_DATASETS[datasetId].data || [];")
    add("                CURRENT_DATASET = datasetId;")
    add("                console.log('[fix.py] switch FIXPY_DATASET:',")
    add("                            datasetId, '->', RAW_DATA.length, 'cau');")
    add("                return true;")
    add("            }")
    add("            return origSwitch.apply(this, arguments);")
    add("        };")
    add("        window.__fixPySwitchPatched = true;")
    add("    }")
    add("    patchSwitchRawData();")
    add("")

    # Patch markCurrentDatasetActive
    add("    function patchMarkActive() {")
    add("        if (window.__fixPyMarkPatched) return;")
    add("        window.markCurrentDatasetActive = function() {")
    add("            var cur = (typeof CURRENT_DATASET !== 'undefined')")
    add("                      ? CURRENT_DATASET : 'tonghop';")
    add("            document.querySelectorAll('.ds-btn, .ds-sub-btn').forEach(function(b) {")
    add("                b.classList.remove('active');")
    add("            });")
    add("            if (cur === 'tonghop') {")
    add("                var tonghopBtn = document.querySelector('.ds-btn[data-dataset=\"tonghop\"]');")
    add("                if (tonghopBtn) tonghopBtn.classList.add('active');")
    add("                return;")
    add("            }")
    add("            if (window.FIXPY_DATASETS && window.FIXPY_DATASETS[cur]) {")
    add("                var fxBtn = document.querySelector('.ds-btn[data-dataset=\"' + cur + '\"]');")
    add("                if (fxBtn) fxBtn.classList.add('active');")
    add("                return;")
    add("            }")
    add("            var subBtn = document.querySelector('.ds-sub-btn[data-dataset=\"' + cur + '\"]');")
    add("            if (subBtn) subBtn.classList.add('active');")
    add("            var cnBtn = document.querySelector('.ds-btn[data-dataset-group=\"chuyen-nganh\"]');")
    add("            if (cnBtn && subBtn) cnBtn.classList.add('active');")
    add("        };")
    add("        window.__fixPyMarkPatched = true;")
    add("    }")
    add("")

    # Patch getLimitedData
    add("    function patchGetLimitedData() {")
    add("        if (window.__fixPyLimitedPatched) return;")
    add("        var origGet = window.getLimitedData")
    add("                    || (typeof getLimitedData !== 'undefined' ? getLimitedData : null);")
    add("        if (typeof origGet !== 'function') return;")
    add("        window.getLimitedData = function() {")
    add("            var currentDs = (typeof CURRENT_DATASET !== 'undefined')")
    add("                            ? CURRENT_DATASET : 'tonghop';")
    add("            if (currentDs === 'tonghop') {")
    add("                return origGet.apply(this, arguments);")
    add("            }")
    add("            if (!window.FIXPY_DATASETS || !window.FIXPY_DATASETS[currentDs]) {")
    add("                return origGet.apply(this, arguments);")
    add("            }")
    add("            var info = (typeof getTierInfo === 'function')")
    add("                       ? getTierInfo() : {};")
    add("            if (info.tier === 'active') {")
    add("                return RAW_DATA;")
    add("            }")
    add("            var override = window.__onboardingOverride;")
    add("            if (override && Array.isArray(override) && override.length > 0")
    add("                && typeof state !== 'undefined' && state")
    add("                && !state.search && !state.hsk && !state.subject) {")
    add("                var currentStts = {};")
    add("                RAW_DATA.forEach(function(r) { currentStts[r.stt] = true; });")
    add("                var filtered = override.filter(function(r) {")
    add("                    return currentStts[r.stt];")
    add("                });")
    add("                if (filtered.length > 0) {")
    add("                    var max = info.maxQuestions || 60;")
    add("                    return filtered.slice(0, max);")
    add("                }")
    add("            }")
    add("            var max2 = info.maxQuestions || 60;")
    add("            return RAW_DATA.slice(0, max2);")
    add("        };")
    add("        window.__fixPyLimitedPatched = true;")
    add("    }")
    add("")

    # bindTab
    add("    function bindTab(dsId) {")
    add("        var btn = document.querySelector('.ds-btn[data-dataset=\"' + dsId + '\"]');")
    add("        if (!btn) return;")
    add("        if (btn.__fixPyBound) return;")
    add("        var cloned = btn.cloneNode(true);")
    add("        btn.parentNode.replaceChild(cloned, btn);")
    add("        cloned.__fixPyBound = true;")
    add("        cloned.addEventListener('click', function(e) {")
    add("            e.stopImmediatePropagation();")
    add("            e.stopPropagation();")
    add("            e.preventDefault();")
    add("            console.log('[fix.py] click tab ' + dsId);")
    add("            var sub = document.getElementById('dsSubWrap');")
    add("            if (sub) sub.style.display = 'none';")
    add("            document.querySelectorAll('.ds-btn, .ds-sub-btn').forEach(function(b) {")
    add("                b.classList.remove('active');")
    add("            });")
    add("            this.classList.add('active');")
    add("            if (typeof window.__switchRawData === 'function') {")
    add("                window.__switchRawData(dsId);")
    add("            }")
    add("            if (typeof state !== 'undefined' && state) {")
    add("                state.search = '';")
    add("                state.hsk = '';")
    add("                state.subject = '';")
    add("            }")
    add("            try {")
    add("                var si = document.getElementById('searchInput');")
    add("                var hf = document.getElementById('hskFilter');")
    add("                var sf = document.getElementById('subjectFilter');")
    add("                if (si) si.value = '';")
    add("                if (hf) hf.value = '';")
    add("                if (sf) sf.value = '';")
    add("                var cb = document.getElementById('clearSearchBtn');")
    add("                if (cb) cb.classList.remove('show');")
    add("            } catch(err) {}")
    add("            var savedTopics = null;")
    add("            if (typeof loadOnboardingSelection === 'function') {")
    add("                try {")
    add("                    var saved = loadOnboardingSelection();")
    add("                    if (saved && saved.topics && saved.topics.length > 0) {")
    add("                        savedTopics = saved.topics;")
    add("                    }")
    add("                } catch(e) {}")
    add("            }")
    add("            try {")
    add("                if (typeof buildFilters === 'function') buildFilters();")
    add("                if (typeof applyFilter === 'function') applyFilter();")
    add("                if (typeof updateResultCount === 'function') updateResultCount();")
    add("                if (savedTopics && savedTopics.length > 0")
    add("                    && typeof applyOnboardingSelection === 'function') {")
    add("                    try {")
    add("                        var cfg = (typeof getOnboardingConfig === 'function')")
    add("                                  ? getOnboardingConfig() : null;")
    add("                        if (cfg) {")
    add("                            var saved2 = loadOnboardingSelection();")
    add("                            window.__onboardingAutoPicked = saved2 ? !!saved2.auto_picked : false;")
    add("                            applyOnboardingSelection(savedTopics, false);")
    add("                        }")
    add("                    } catch(e2) {")
    add("                        console.warn('[fix.py] applyOnboarding error:', e2);")
    add("                    }")
    add("                }")
    add("                document.querySelectorAll('.ds-btn, .ds-sub-btn').forEach(function(b) {")
    add("                    b.classList.remove('active');")
    add("                });")
    add("                this.classList.add('active');")
    add("            } catch(err) {")
    add("                console.warn('[fix.py] re-render error:', err);")
    add("            }")
    add("            setTimeout(function() {")
    add("                var mainEl = document.getElementById('mainContent');")
    add("                if (mainEl) {")
    add("                    var yOffset = mainEl.getBoundingClientRect().top")
    add("                                + window.scrollY - 100;")
    add("                    window.scrollTo({ top: yOffset, behavior: 'smooth' });")
    add("                }")
    add("            }, 100);")
    add("        }, true);")
    add("    }")
    add("")

    # patchApplyOnboarding
    add("    function patchApplyOnboarding() {")
    add("        if (window.__fixPyApplyOnbPatched) return;")
    add("        var origApply = window.applyOnboardingSelection")
    add("                      || (typeof applyOnboardingSelection !== 'undefined'")
    add("                          ? applyOnboardingSelection : null);")
    add("        if (typeof origApply !== 'function') return;")
    add("        window.applyOnboardingSelection = function(topics, scrollTop) {")
    add("            var currentDs = (typeof CURRENT_DATASET !== 'undefined')")
    add("                            ? CURRENT_DATASET : 'tonghop';")
    add("            if (!window.FIXPY_DATASETS || !window.FIXPY_DATASETS[currentDs]) {")
    add("                return origApply.apply(this, arguments);")
    add("            }")
    add("            var cfg = (typeof getOnboardingConfig === 'function')")
    add("                      ? getOnboardingConfig() : null;")
    add("            if (!cfg) return;")
    add("            var maxQ = cfg.max_questions;")
    add("            var isUnlimitedQ = (maxQ === -1 || maxQ === Infinity);")
    add("            var allowedHsk;")
    add("            if (cfg.hsk_allowed && Array.isArray(cfg.hsk_allowed) && cfg.hsk_allowed.length > 0) {")
    add("                allowedHsk = cfg.hsk_allowed.map(function(n) { return 'HSK' + n; });")
    add("            } else {")
    add("                allowedHsk = (typeof getAllowedHskList === 'function')")
    add("                             ? getAllowedHskList()")
    add("                             : ['HSK1','HSK2','HSK3','HSK4','HSK5','HSK6'];")
    add("            }")
    add("            var pool = RAW_DATA.filter(function(r) {")
    add("                if (allowedHsk.indexOf(r.hsk) === -1) return false;")
    add("                var s = (r.subject || '').trim();")
    add("                return topics.indexOf(s) !== -1;")
    add("            });")
    add("            pool.sort(function(a, b) {")
    add("                return (parseInt(a.stt) || 0) - (parseInt(b.stt) || 0);")
    add("            });")
    add("            var final = [];")
    add("            if (isUnlimitedQ) {")
    add("                final = pool.slice();")
    add("            } else {")
    add("                var maxPerTopic = (typeof getMaxPerTopic === 'function')")
    add("                                 ? getMaxPerTopic(maxQ, RAW_DATA)")
    add("                                 : Math.max(1, Math.ceil(maxQ / Math.max(1, topics.length)));")
    add("                var topicCount = {};")
    add("                var perHsk = Math.ceil(maxQ / allowedHsk.length);")
    add("                var hskCount = {};")
    add("                allowedHsk.forEach(function(h) { hskCount[h] = 0; });")
    add("                for (var i = 0; i < pool.length && final.length < maxQ; i++) {")
    add("                    var r = pool[i];")
    add("                    var s = (r.subject || '').trim() || '__no_subject__';")
    add("                    if ((topicCount[s] || 0) >= maxPerTopic) continue;")
    add("                    if (r.hsk && hskCount[r.hsk] !== undefined && hskCount[r.hsk] >= perHsk) continue;")
    add("                    final.push(r);")
    add("                    topicCount[s] = (topicCount[s] || 0) + 1;")
    add("                    if (r.hsk && hskCount[r.hsk] !== undefined) hskCount[r.hsk]++;")
    add("                }")
    add("                if (final.length < maxQ) {")
    add("                    var usedIds = {};")
    add("                    final.forEach(function(r) { usedIds[r.stt] = true; });")
    add("                    for (var p = 0; p < pool.length && final.length < maxQ; p++) {")
    add("                        var rp = pool[p];")
    add("                        if (usedIds[rp.stt]) continue;")
    add("                        var sp = (rp.subject || '').trim() || '__no_subject__';")
    add("                        if ((topicCount[sp] || 0) >= maxPerTopic) continue;")
    add("                        final.push(rp);")
    add("                        usedIds[rp.stt] = true;")
    add("                        topicCount[sp] = (topicCount[sp] || 0) + 1;")
    add("                    }")
    add("                }")
    add("                if (final.length > maxQ) final = final.slice(0, maxQ);")
    add("            }")
    add("            final.sort(function(a, b) {")
    add("                return (parseInt(a.stt) || 0) - (parseInt(b.stt) || 0);")
    add("            });")
    add("            window.__onboardingOverride = final;")
    add("            if (typeof state !== 'undefined' && state) {")
    add("                state.search = '';")
    add("                state.hsk = '';")
    add("                state.subject = '';")
    add("            }")
    add("            try {")
    add("                var si = document.getElementById('searchInput');")
    add("                if (si) si.value = '';")
    add("                var cb = document.getElementById('clearSearchBtn');")
    add("                if (cb) cb.classList.remove('show');")
    add("            } catch(e) {}")
    add("            if (typeof applyFilter === 'function') applyFilter();")
    add("            if (typeof updateResultCount === 'function') updateResultCount();")
    add("            if (typeof showOnboardingActiveBanner === 'function') {")
    add("                showOnboardingActiveBanner(topics, final.length);")
    add("            }")
    add("            if (scrollTop) {")
    add("                setTimeout(function() {")
    add("                    var mainEl = document.getElementById('mainContent');")
    add("                    if (mainEl) {")
    add("                        var yOffset = mainEl.getBoundingClientRect().top")
    add("                                    + window.scrollY - 100;")
    add("                        window.scrollTo({ top: yOffset, behavior: 'smooth' });")
    add("                    }")
    add("                }, 200);")
    add("            }")
    add("        };")
    add("        window.__fixPyApplyOnbPatched = true;")
    add("    }")
    add("")

    # INIT
    add("    function bindAll() {")
    add("        patchSwitchRawData();")
    add("        patchGetLimitedData();")
    add("        patchApplyOnboarding();")
    add("        NEW_IDS.forEach(bindTab);")
    add("        patchMarkActive();")
    add("    }")
    add("")
    add("    if (document.readyState === 'loading') {")
    add("        document.addEventListener('DOMContentLoaded', bindAll);")
    add("    } else {")
    add("        bindAll();")
    add("    }")
    add("")
    add("    var _timer = null;")
    add("    var observer = new MutationObserver(function() {")
    add("        clearTimeout(_timer);")
    add("        _timer = setTimeout(bindAll, 200);")
    add("    });")
    add("    if (document.body) {")
    add("        observer.observe(document.body, { childList: true, subtree: true });")
    add("    }")
    add("")
    add("    console.log('[fix.py] Da bind ' + NEW_IDS.length + ' tab:', NEW_IDS);")
    add("})();")
    add("</script>")

    return "\n".join(L)


# =================================================================
#  MAIN
# =================================================================
def main():
    print("=" * 62)
    print("[fix.py] Auto-scan data/ -> them tab rieng cho moi file")
    print("=" * 62)

    if not os.path.isfile(INDEX_HTML):
        print("[X] Khong thay " + INDEX_HTML + ". Chay convert.py truoc.")
        sys.exit(1)

    with open(INDEX_HTML, "r", encoding="utf-8") as f:
        html = f.read()

    # ═══════════════════════════════════════════════════════════════
    #  SCAN tab thường từ data/
    # ═══════════════════════════════════════════════════════════════
    datasets = scan_data_dir()

    # ═══════════════════════════════════════════════════════════════
    #  ⭐ ĐỌC FILE TỪ VỰNG PREMIUM
    # ═══════════════════════════════════════════════════════════════
    vocab_data = []
    vocab_real_path = None
    if HAS_VOCAB_MODULE:
        print("")
        print("[VOCAB] Kiem tra file tu vung: " + VOCAB_FILE)
        vocab_real_path = _find_file_safe(VOCAB_FILE)
        if vocab_real_path:
            print("[VOCAB] Tim thay: " + os.path.basename(vocab_real_path))
            vocab_data = read_vocab_excel(vocab_real_path)
            if not vocab_data:
                print("[VOCAB] [!] File rong hoac loi")
        else:
            print("[VOCAB] Khong co file tu vung - bo qua")

    # ═══════════════════════════════════════════════════════════════
    #  KIỂM TRA CÓ GÌ MỚI CẦN PATCH KHÔNG
    # ═══════════════════════════════════════════════════════════════
    all_new = []
    for ds in datasets:
        marker = 'data-dataset="' + ds["id"] + '"'
        if marker not in html:
            all_new.append(ds)
        else:
            print("   [skip] '" + ds["id"] + "' da co trong HTML")

    # ⭐ Kiểm tra tab từ vựng đã có chưa
    vocab_exists_in_html = 'data-dataset="' + VOCAB_ID + '"' in html
    add_vocab = bool(vocab_data) and not vocab_exists_in_html

    if not all_new and not add_vocab:
        print("")
        print("[fix.py] Tat ca da co - khong can patch.")
        return

    print("")
    print("[fix.py] Se them:")
    for ds in all_new:
        print("   - [tab] " + ds["name"] + " (" + str(ds["count"]) + " cau)")
    if add_vocab:
        print("   - [👑 PREMIUM] " + VOCAB_LABEL + " (" + str(len(vocab_data)) + " tu)")

    # ═══════════════════════════════════════════════════════════════
    #  PATCH 2: BUTTONS
    # ═══════════════════════════════════════════════════════════════
    print("")
    print("[PATCH 2] Them button tabs...")
    new_btns = ""
    for ds in all_new:
        label = ds["name"]
        new_btns += (
            '\n        <button class="ds-btn ds-btn-primary" '
            'data-dataset="' + ds["id"] + '">\n'
            '            <i class="fas ' + ds["icon"] + '"></i>\n'
            '            <span>' + _js_str(label) + '</span>\n'
            '        </button>'
        )

    # ⭐ Thêm tab Từ vựng Premium
    if add_vocab:
        new_btns += build_vocab_tab_html(VOCAB_ID, VOCAB_LABEL)

    # Chèn SAU nút "Tổng hợp"
    pat_after_tonghop = re.compile(
        r'(<button[^>]*class="[^"]*ds-btn[^"]*"[^>]*data-dataset="tonghop"[^>]*>.*?</button>)',
        re.MULTILINE | re.DOTALL
    )
    html, n = pat_after_tonghop.subn(r'\1' + new_btns, html, count=1)

    if n > 0:
        print("   [OK] Da chen button (sau 'Tong hop')")
    else:
        print("   [!] Khong thay nut 'tonghop' -> fallback truoc 'chuyen-nganh'")
        pat_before_cn = re.compile(
            r'(\s*)(<button\s+class="[^"]*ds-btn[^"]*"\s+[^>]*data-dataset-group="chuyen-nganh")',
            re.MULTILINE
        )
        html, n = pat_before_cn.subn(r'\1' + new_btns + '\n        ' + r'\2', html, count=1)
        if n == 0:
            print("[X] Khong tim thay ca nut 'tonghop' lan 'chuyen-nganh'")
            sys.exit(1)
        print("   [OK] Da chen button (fallback)")

    # ═══════════════════════════════════════════════════════════════
    #  PATCH 3: CSS
    # ═══════════════════════════════════════════════════════════════
    print("")
    print("[PATCH 3] CSS layout...")

    css_lines = []
    css_lines.append("")
    css_lines.append("/* ==== FIX.PY: AUTO-FIT LAYOUT ==== */")
    css_lines.append("@media (max-width: 768px) {")
    css_lines.append("    .ds-main-row {")
    css_lines.append("        grid-template-columns: repeat(2, minmax(0, 1fr)) !important;")
    css_lines.append("        gap: .5rem !important;")
    css_lines.append("    }")
    css_lines.append("}")
    css_lines.append("@media (min-width: 769px) {")
    css_lines.append("    .ds-main-row {")
    css_lines.append("        grid-template-columns: repeat(4, minmax(0, 1fr)) !important;")
    css_lines.append("        gap: .55rem !important;")
    css_lines.append("    }")
    css_lines.append("}")

    # Grid cho tab con chuyên ngành
    css_lines.append("")
    css_lines.append("/* ==== GRID ĐỀU CHO TAB CON CHUYÊN NGÀNH ==== */")
    css_lines.append(".ds-sub-grid {")
    css_lines.append("    display: grid !important;")
    css_lines.append("    grid-template-columns: repeat(2, minmax(0, 1fr)) !important;")
    css_lines.append("    gap: .55rem !important;")
    css_lines.append("    align-items: stretch !important;")
    css_lines.append("}")
    css_lines.append("@media (min-width: 600px) {")
    css_lines.append("    .ds-sub-grid { grid-template-columns: repeat(3, minmax(0, 1fr)) !important; }")
    css_lines.append("}")
    css_lines.append("@media (min-width: 900px) {")
    css_lines.append("    .ds-sub-grid { grid-template-columns: repeat(4, minmax(0, 1fr)) !important; }")
    css_lines.append("}")
    css_lines.append(".ds-sub-btn {")
    css_lines.append("    width: 100% !important;")
    css_lines.append("    min-width: 0 !important;")
    css_lines.append("    max-width: 100% !important;")
    css_lines.append("    height: 100% !important;")
    css_lines.append("    min-height: 44px !important;")
    css_lines.append("    padding: .55rem .7rem !important;")
    css_lines.append("    justify-content: flex-start !important;")
    css_lines.append("    border-radius: 12px !important;")
    css_lines.append("    font-size: clamp(.68rem, 1.9vw, .82rem) !important;")
    css_lines.append("    white-space: normal !important;")
    css_lines.append("    word-break: break-word !important;")
    css_lines.append("    overflow-wrap: anywhere !important;")
    css_lines.append("    line-height: 1.25 !important;")
    css_lines.append("    text-align: left !important;")
    css_lines.append("}")
    css_lines.append(".ds-sub-btn span {")
    css_lines.append("    flex: 1 1 auto !important;")
    css_lines.append("    min-width: 0 !important;")
    css_lines.append("    display: -webkit-box !important;")
    css_lines.append("    -webkit-line-clamp: 2 !important;")
    css_lines.append("    -webkit-box-orient: vertical !important;")
    css_lines.append("    overflow: hidden !important;")
    css_lines.append("}")

    # Force style tab thường giống "Tổng hợp"
    for ds in all_new:
        i = ds["id"]
        sel = '.ds-btn[data-dataset="' + i + '"]'
        css_lines.append("")
        css_lines.append("/* Style cho tab " + i + " */")
        css_lines.append(sel + " {")
        css_lines.append("    background: var(--surface) !important;")
        css_lines.append("    border-color: var(--border) !important;")
        css_lines.append("    color: var(--text) !important;")
        css_lines.append("}")
        css_lines.append(sel + " i:first-child { color: var(--primary) !important; }")
        css_lines.append(sel + ":hover {")
        css_lines.append("    border-color: var(--primary) !important;")
        css_lines.append("    background: var(--primary-light) !important;")
        css_lines.append("}")
        css_lines.append(sel + ".active {")
        css_lines.append("    background: linear-gradient(135deg, #4f46e5, #7c3aed) !important;")
        css_lines.append("    color: #fff !important;")
        css_lines.append("    border-color: transparent !important;")
        css_lines.append("    box-shadow: 0 4px 12px rgba(124, 58, 237, .35) !important;")
        css_lines.append("}")
        css_lines.append(sel + ".active i:first-child { color: #fff !important; }")
        css_lines.append('[data-theme="dark"] ' + sel + " {")
        css_lines.append("    background: var(--surface) !important;")
        css_lines.append("    border-color: var(--border) !important;")
        css_lines.append("}")
        css_lines.append('[data-theme="dark"] ' + sel + ".active {")
        css_lines.append("    background: linear-gradient(135deg, #4f46e5, #7c3aed) !important;")
        css_lines.append("    border-color: transparent !important;")
        css_lines.append("}")

    # ⭐ CSS cho tab Từ vựng Premium
    if add_vocab:
        css_lines.append("")
        css_lines.append("/* ==== VOCAB PREMIUM CSS ==== */")
        css_lines.append(build_vocab_css(VOCAB_ID))

    css = "\n".join(css_lines) + "\n"

    pat_style = re.compile(r'(\s*)(</style>)', re.MULTILINE)
    html, n = pat_style.subn(r'\1' + css + r'\1\2', html, count=1)
    if n == 0:
        print("   [!] Khong tim thay </style>")
    else:
        print("   [OK] Da inject CSS")

    # ═══════════════════════════════════════════════════════════════
    #  PATCH 4: JS
    # ═══════════════════════════════════════════════════════════════
    print("")
    print("[PATCH 4] JS binding...")

    # IDs cho tab thường
    ids_js = json.dumps([ds["id"] for ds in all_new])

    # Dict datasets cho JS
    datasets_dict = {}
    for ds in all_new:
        datasets_dict[ds["id"]] = ds

    # ⭐ THÊM VOCAB VÀO FIXPY_DATASETS
    if add_vocab:
        datasets_dict[VOCAB_ID] = {
            "id": VOCAB_ID,
            "name": VOCAB_LABEL,
            "icon": "fa-book",
            "color": "#f59e0b",
            "data": vocab_data,
            "count": len(vocab_data),
            "source": os.path.basename(vocab_real_path),
            "type": "premium",
            "group": "fixpy",
        }

    datasets_json = _escape_json_for_script(datasets_dict)
    js = build_js_override(ids_js, datasets_json)

    # ⭐ JS riêng cho module Từ vựng Premium (bind lock + inject block)
    if add_vocab:
        js += build_vocab_js_override(VOCAB_ID)

    # ⭐ Modal upgrade HTML
    modal_html = ""
    if add_vocab:
        modal_html = build_vocab_modal_html()

    pat_body = re.compile(r'(\s*)(</body>)', re.MULTILINE)
    html, n = pat_body.subn(r'\1' + modal_html + '\n' + js + r'\1\2', html, count=1)
    if n == 0:
        print("[X] Khong tim thay </body>")
        sys.exit(1)
    print("   [OK] Da inject JS + modal")

    # GHI FILE
    with open(INDEX_HTML, "w", encoding="utf-8") as f:
        f.write(html)

    size_kb = os.path.getsize(INDEX_HTML) / 1024

    print("")
    print("=" * 62)
    print("[fix.py] HOAN TAT! Da patch " + INDEX_HTML)
    print("[fix.py] Kich thuoc: " + str(round(size_kb, 1)) + " KB")

    if all_new:
        print("[fix.py] Tab thuong da them:")
        for ds in all_new:
            print("   - " + ds["name"] + " (" + str(ds["count"]) + " cau)")

    if add_vocab:
        print("[fix.py] 👑 Tab Tu vung PREMIUM: " + str(len(vocab_data)) + " tu")
        print("[fix.py]    (chi Admin + Premium moi mo duoc)")

    print("[fix.py] Layout: PC 4 cot - Mobile 2 cot")
    print("[fix.py] CHI 1 TAB ACTIVE tai mot thoi diem")
    print("[fix.py] Dropdown 'Chuyen nganh' KHONG hien thi tab fix.py")
    print("=" * 62)


if __name__ == "__main__":
    main()
